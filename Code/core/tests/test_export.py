import csv
import io
import json

from openpyxl import load_workbook

from docintel.export import NOT_FOUND, pictures_zip, to_csv, to_json, to_json_many, to_xlsx
from sdk_fixtures import invoice_fields, record, string


def two_records():
    a = record(edits={"InvoiceId": {"value": "KA-0043", "at": "t", "by": "User"}})
    fields = invoice_fields()
    fields["VendorName"] = string("=HYPERLINK(\"http://x\")")
    b = record(fields, id="fedcba9876543210", filename="second.pdf", selection=["VendorName", "AmountDue", "Items"])
    return [a, b]


def test_xlsx_sheets_columns_and_values():
    wb = load_workbook(io.BytesIO(to_xlsx(two_records(), 0.80)))
    assert wb.sheetnames == ["Documents", "Line items", "Details"]
    docs = list(wb["Documents"].iter_rows(values_only=True))
    header = docs[0]
    assert header[:5] == ("Document", "Document ID", "Model", "Pages", "Status")
    assert "Invoice total" in header and "Amount due" in header
    assert header[-4:] == ("Fields needing review", "Edited fields", "Pictures", "Problem")
    first = dict(zip(header, docs[1]))
    assert first["Invoice number"] == "KA-0043"  # the corrected value
    assert first["Edited fields"] == "Invoice number"
    assert first["Invoice total"] == 1180.0  # numbers stay numbers
    assert first["Due date"] == NOT_FOUND
    assert first["Invoice date"] == "2026-09-14"
    second = dict(zip(header, docs[2]))
    assert second["Vendor name"].startswith("'=")  # formula neutralised
    assert second["Amount due"] == NOT_FOUND and second["Invoice total"] is None  # not selected for this one

    lines = list(wb["Line items"].iter_rows(values_only=True))
    assert lines[0][:3] == ("Document", "Table", "Row") and lines[0][-1] == "Edited"
    assert [r[0] for r in lines[1:]] == ["invoice.pdf", "invoice.pdf", "second.pdf", "second.pdf"]
    item_rows = [r for r in lines[1:] if r[1] == "Line items"]
    assert len(item_rows) == 4 and item_rows[0][2] == 1

    details = list(wb["Details"].iter_rows(values_only=True))
    assert details[0] == ("Document", "Field", "Label", "Value", "Raw text", "Confidence", "Needs review", "Edited",
                          "Model value", "Page")
    inv = next(r for r in details if r[1] == "InvoiceId" and r[0] == "invoice.pdf")
    assert inv[3] == "KA-0043" and inv[7] == "Yes" and inv[8] == "KA-0042" and inv[5] == 0.72
    cell = next(r for r in details if r[1] == "Items[1].Amount")
    assert cell[6] == "Yes"  # 0.55 needs review


def test_csv_one_row_per_document_with_confidence_columns():
    text = to_csv(two_records(), 0.80).decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    header = rows[0]
    assert "Invoice total" in header and "Invoice total confidence" in header
    assert len(rows) == 3
    first = dict(zip(header, rows[1]))
    assert first["Invoice number"] == "KA-0043" and first["Invoice number confidence"] == "0.72"
    assert first["Due date"] == NOT_FOUND and first["Line items"] == "2"
    assert first["Edited fields"] == "Invoice number"


def test_failed_and_unfinished_documents_are_listed_with_their_status():
    failed = record(id="1111111111111111", filename="bad.pdf", status="failed", analysis=None,
                    error={"code": "unsupported", "message": "Could not read this file."})
    queued = record(id="2222222222222222", filename="later.pdf", status="queued", analysis=None)
    rows = list(csv.reader(io.StringIO(to_csv([two_records()[0], failed, queued], 0.8).decode("utf-8-sig"))))
    header = rows[0]
    by_doc = {r[0]: dict(zip(header, r)) for r in rows[1:]}
    assert by_doc["invoice.pdf"]["Status"] == "Done"
    assert by_doc["bad.pdf"]["Status"] == "Failed" and by_doc["bad.pdf"]["Problem"] == "Could not read this file."
    assert by_doc["later.pdf"]["Status"] == "Not finished"
    wb = load_workbook(io.BytesIO(to_xlsx([two_records()[0], failed], 0.8)))
    docs = list(wb["Documents"].iter_rows(values_only=True))
    assert [r[4] for r in docs[1:]] == ["Done", "Failed"]
    assert "Pictures" not in wb.sheetnames  # no document asked for pictures


def test_pictures_are_listed_and_zipped():
    import zipfile

    rec = record(include_images=True, figures=[{"id": "1.1", "page": 1, "polygon": [0, 0, 1, 0, 1, 1, 0, 1],
                                                "caption": "", "size": 3}])
    wb = load_workbook(io.BytesIO(to_xlsx([rec], 0.8)))
    rows = list(wb["Pictures"].iter_rows(values_only=True))
    assert rows[1][:3] == ("invoice.pdf", "1.1", 1) and rows[1][4] == "invoice-012345/picture-1.1.png"
    assert dict(zip(*list(wb["Documents"].iter_rows(values_only=True))[:2]))["Pictures"] == 1
    data = pictures_zip([rec], lambda r, fid: b"png")
    assert zipfile.ZipFile(io.BytesIO(data)).namelist() == ["invoice-012345/picture-1.1.png"]
    assert pictures_zip([record()], lambda r, fid: b"") is None
    doc = json.loads(to_json(rec, 0.8))
    assert doc["pictures"][0]["url"] == "/api/documents/0123456789abcdef/figures/1.1"


def test_json_keeps_full_analysis_and_selected_view():
    doc = json.loads(to_json(two_records()[0], 0.80))
    assert doc["api_version"] == "2024-11-30" and doc["model"] == "prebuilt-invoice"
    assert set(doc["fields"]) == {"VendorName", "InvoiceId", "InvoiceDate", "DueDate", "VendorAddress",
                                  "CurrencyCode", "SubTotal", "TotalTax", "InvoiceTotal"}
    assert doc["fields"]["DueDate"] == {**doc["fields"]["DueDate"], "found": False, "value": None}
    assert doc["fields"]["InvoiceId"]["edited"] and doc["fields"]["InvoiceId"]["model_value"] == "KA-0042"
    assert len(doc["tables"]["Items"]) == 2
    assert "VendorName" in doc["analysis"]["documents"][0]["fields"]  # every field the service returned is kept
    many = json.loads(to_json_many(two_records(), 0.8))
    assert len(many["documents"]) == 2
