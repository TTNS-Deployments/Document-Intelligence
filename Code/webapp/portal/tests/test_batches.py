import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from openpyxl import load_workbook

from docintel import batches, results

from .base import WebTestCase


class BatchPageTests(WebTestCase):
    def bulk(self, *sample_ids, name="October invoices", fields=("VendorName", "PurchaseOrder", "InvoiceTotal")):
        """Upload like the browser does: one file per request, then mark the upload finished."""
        batch_id = ""
        for i, sid in enumerate(sample_ids):
            form = {"model": "prebuilt-invoice", "fields": list(fields),
                    "files": [SimpleUploadedFile(f"{sid}.pdf", self.sample_bytes(sid))]}
            if batch_id:
                form["batch"] = batch_id
            else:
                form.update(batch_size=str(len(sample_ids)), batch_name=name)
            batch_id = self.client.post("/api/documents", form).json()["batch_id"]
        self.assertEqual(self.client.post(f"/api/batches/{batch_id}/finish").json()["expected"], len(sample_ids))
        return batch_id

    def test_bulk_upload_becomes_a_named_batch_on_the_batches_page(self):
        self.sign_in("Admin")
        batch_id = self.bulk("inv-01-gst-software", "inv-02-us-logistics", "inv-03-uk-interiors")
        self.assertEqual(batches.get(batch_id)["name"], "October invoices")
        self.wait()
        page = self.client.get("/batches")
        self.assertContains(page, "October invoices")
        self.assertContains(page, "3 / 3")
        self.assertContains(page, f"/api/export.xlsx?batch={batch_id}")
        self.assertNotContains(page, 'hx-trigger="every 5s"')  # nothing running: no polling

    def test_single_uploads_and_samples_are_not_batches(self):
        self.sign_in()
        self.upload(("a.pdf", self.sample_bytes("inv-02-us-logistics")))
        self.client.post("/api/samples/inv-03-uk-interiors")
        self.wait()
        self.assertContains(self.client.get("/batches"), "No batches yet")

    def test_batch_page_reviews_documents_together_and_exports_corrections(self):
        self.sign_in("Admin")
        batch_id = self.bulk("inv-02-us-logistics", "inv-01-gst-software")
        self.wait()
        page = self.client.get(f"/batches/{batch_id}")
        self.assertContains(page, "October invoices")
        docs = results.batch(batch_id)
        to_review = next(d for d in docs if d["review_count"])
        self.assertContains(page, f'data-first="{to_review["id"]}"')  # opens the first document needing review
        self.assertContains(page, "Next needing review")
        self.assertContains(page, "All 2 documents are finished")
        # a correction made while reviewing goes into the batch Excel file
        self.client.post(f"/documents/{to_review['id']}/edit", {"path": "PurchaseOrder", "value": "PO-CHECKED"})
        partial = self.client.get(f"/partials/batches/{batch_id}/documents?current={to_review['id']}")
        self.assertContains(partial, "Checked")
        self.assertContains(partial, 'hx-swap-oob="true"')
        wb = load_workbook(io.BytesIO(self.client.get(f"/api/export.xlsx?batch={batch_id}").content))
        rows = list(wb["Documents"].iter_rows(values_only=True))
        by_doc = {r[0]: dict(zip(rows[0], r)) for r in rows[1:]}
        self.assertEqual(by_doc["inv-01-gst-software.pdf"]["Purchase order"], "PO-CHECKED")
        self.assertEqual(len(by_doc), 2)

    def test_batch_still_running_has_no_excel_and_polls(self):
        from docintel import jobs

        self.sign_in("Admin")
        batch_id = self.bulk("inv-02-us-logistics", "inv-03-uk-interiors")
        self.wait()
        jobs.submit_file("late.pdf", self.sample_bytes("inv-04-eu-energy"), model_id="prebuilt-invoice",
                         fields=["VendorName"], owner="x", owner_role="Admin", batch_id=batch_id, start=False)
        page = self.client.get(f"/batches/{batch_id}")
        self.assertContains(page, "Waiting")
        self.assertNotContains(page, "data-batch-excel")
        self.assertContains(page, 'hx-trigger="every 3s"')
        self.assertContains(self.client.get("/batches"), 'hx-trigger="every 5s"')

    def test_rename_is_escaped_and_validated(self):
        self.sign_in()
        batch_id = self.bulk("inv-02-us-logistics", "inv-03-uk-interiors")
        r = self.client.post(f"/api/batches/{batch_id}/rename", {"name": "<b>Q4</b>"}, HTTP_HX_REQUEST="true")
        self.assertEqual(r.content.decode(), "&lt;b&gt;Q4&lt;/b&gt;")
        self.assertEqual(self.client.post(f"/api/batches/{batch_id}/rename", {"name": " "}).status_code, 400)
        self.assertEqual(self.client.post("/api/batches/0123456789abcdef/rename", {"name": "x"}).status_code, 404)

    def test_only_admins_delete_batches(self):
        self.sign_in("User")
        batch_id = self.bulk("inv-02-us-logistics", "inv-03-uk-interiors")
        self.wait()
        self.assertEqual(self.client.post(f"/api/batches/{batch_id}/delete").status_code, 403)
        self.assertNotContains(self.client.get(f"/batches/{batch_id}"), "Delete batch")
        admin = Client()
        self.sign_in("Admin", client=admin)
        self.assertContains(admin.get(f"/batches/{batch_id}"), "Delete batch")
        r = admin.post(f"/api/batches/{batch_id}/delete")
        self.assertEqual(r.json()["documents"], 2)
        self.assertEqual(admin.get(f"/batches/{batch_id}").status_code, 404)
        self.assertEqual(results.batch(batch_id), [])

    def test_unknown_batch(self):
        self.sign_in()
        self.assertEqual(self.client.get("/batches/0123456789abcdef").status_code, 404)
        self.assertEqual(self.client.post("/api/batches/0123456789abcdef/finish").status_code, 404)
