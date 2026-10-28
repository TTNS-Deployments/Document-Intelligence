import io
import os

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings

from docintel import results
from docintel.config import reset_settings_cache

from .base import WebTestCase


def pdf(pages=1):
    from pypdf import PdfWriter

    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=612, height=792)
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


class UploadValidationTests(WebTestCase):
    def setUp(self):
        super().setUp()
        self.sign_in()

    def assertRejected(self, response, code):
        self.assertEqual(response.status_code, 400, response.content)
        body = response.json()
        self.assertEqual(body["accepted"], [])
        self.assertEqual(body["rejected"][0]["code"], code)
        self.assertEqual(results.list_summaries(), [])  # nothing stored
        return body["rejected"][0]["message"]

    def test_wrong_type(self):
        self.assertIn("only PDF, JPEG and PNG", self.assertRejected(self.upload(("notes.txt", b"hello")), "type"))
        self.assertRejected(self.upload(("archive.zip", b"PK\x03\x04")), "type")

    def test_wrong_magic_bytes(self):
        msg = self.assertRejected(self.upload(("invoice.pdf", b"\x89PNG\r\n\x1a\n" + b"0" * 100)), "content")
        self.assertIn("does not match", msg)
        self.assertRejected(self.upload(("photo.jpg", b"%PDF-1.4 fake")), "content")

    def test_too_large(self):
        os.environ["MAX_UPLOAD_MB"] = "1"
        reset_settings_cache()
        try:
            self.assertRejected(self.upload(("big.pdf", pdf() + b"0" * (1024 * 1024 + 10))), "size")
        finally:
            os.environ.pop("MAX_UPLOAD_MB")
            reset_settings_cache()

    def test_too_many_pages(self):
        self.assertIn("11 pages", self.assertRejected(self.upload(("long.pdf", pdf(11))), "pages"))

    def test_odd_names_are_sanitised(self):
        response = self.upload(("../../evil<script>.pdf", self.sample_bytes()))
        self.assertEqual(response.status_code, 200)
        name = response.json()["accepted"][0]["filename"]
        self.assertNotIn("/", name)
        self.assertNotIn("<", name)
        self.assertTrue(name.endswith(".pdf"))

    def test_batch_with_good_and_bad_files(self):
        response = self.upload(("a.pdf", self.sample_bytes("inv-02-us-logistics")), ("b.txt", b"x"),
                               ("c.png", self.sample_bytes("inv-03-uk-interiors", "png")))
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(len(body["accepted"]), 2)
        self.assertEqual(len(body["rejected"]), 1)
        self.wait()
        batch = self.client.get(f"/partials/batch?batch={body['batch_id']}")
        self.assertContains(batch, "2 of 2 finished")
        self.assertContains(batch, "Download Excel")
        self.assertNotContains(batch, "hx-trigger")  # polling stops when every job has finished

    def test_no_files_or_too_many(self):
        self.assertEqual(self.client.post("/api/documents", {"model": "prebuilt-invoice"}).status_code, 400)
        files = [(f"f{i}.pdf", pdf()) for i in range(11)]
        response = self.upload(*files)
        self.assertEqual(response.status_code, 400)

    def test_unavailable_model_and_unknown_fields(self):
        r = self.upload(("a.pdf", pdf()), model="prebuilt-businessCard", fields=["Emails"])
        self.assertEqual(r.status_code, 400)
        self.assertIn("not available", r.json()["error"])
        r = self.upload(("a.pdf", pdf()), fields=["NoSuchField"])
        self.assertEqual(r.status_code, 400)

    def test_daily_budget_refuses(self):
        os.environ["DAILY_PAGE_LIMIT"] = "2"
        reset_settings_cache()
        r = self.upload(("a.pdf", pdf(2)), ("b.pdf", pdf(1)))
        self.assertEqual(len(r.json()["accepted"]), 1)
        self.assertEqual(r.json()["rejected"][0]["code"], "budget")
        self.assertIn("daily page limit", r.json()["rejected"][0]["message"])

    def test_sample_runs(self):
        r = self.client.post("/api/samples/inv-04-eu-energy", {"fields": ["VendorName", "ServiceStartDate"]})
        self.assertEqual(r.status_code, 200)
        self.wait()
        doc = r.json()["accepted"][0]["id"]
        body = self.client.get(f"/api/documents/{doc}").json()
        self.assertEqual(body["status"], "done")
        self.assertEqual([f["name"] for f in body["view"]["fields"]], ["ServiceStartDate", "VendorName"])
        self.assertEqual(self.client.post("/api/samples/nope").status_code, 404)


class BatchTests(WebTestCase):
    def setUp(self):
        super().setUp()
        self.sign_in()

    def test_files_sent_one_by_one_join_one_batch(self):
        first = self.upload(("a.pdf", self.sample_bytes("inv-02-us-logistics"))).json()
        batch_id = first["batch_id"]
        from django.core.files.uploadedfile import SimpleUploadedFile

        for name, sid in (("b.pdf", "inv-03-uk-interiors"), ("c.pdf", "inv-06-au-consulting")):
            r = self.client.post("/api/documents", {"model": "prebuilt-invoice", "fields": ["VendorName", "InvoiceTotal"],
                                                    "batch": batch_id,
                                                    "files": [SimpleUploadedFile(name, self.sample_bytes(sid))]})
            self.assertEqual(r.json()["batch_id"], batch_id)
        self.wait()
        panel = self.client.get(f"/partials/batch?batch={batch_id}&expected=3")
        self.assertContains(panel, "3 of 3 finished")
        self.assertContains(panel, f"/api/export.xlsx?batch={batch_id}")
        self.assertContains(panel, 'data-complete="1"')

    def test_excel_waits_until_every_document_has_finished(self):
        from docintel import jobs

        batch_id = self.upload(("a.pdf", self.sample_bytes("inv-02-us-logistics"))).json()["batch_id"]
        self.wait()
        jobs.submit_file("late.pdf", self.sample_bytes("inv-04-eu-energy"), model_id="prebuilt-invoice",
                         fields=["VendorName"], owner="x", owner_role="User", batch_id=batch_id, start=False)
        panel = self.client.get(f"/partials/batch?batch={batch_id}")
        self.assertContains(panel, "Available when every document has finished (1 of 2 so far)")
        self.assertNotContains(panel, "data-batch-excel")
        self.assertContains(panel, 'data-complete="0"')
        self.assertContains(panel, "hx-trigger")
        # the browser is still uploading: not complete until the expected number of documents is there
        self.assertContains(self.client.get(f"/partials/batch?batch={batch_id}&expected=5"), "1 of 5 finished")

    def test_batch_excel_lists_every_document_including_failures(self):
        import io

        from openpyxl import load_workbook
        from reportlab.pdfgen import canvas

        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        c.drawString(72, 720, "DOCINTEL-MOCK-FAIL")
        c.save()
        body = self.upload(("good.pdf", self.sample_bytes("inv-02-us-logistics")), ("bad.pdf", buf.getvalue())).json()
        self.wait()
        r = self.client.get(f"/api/export.xlsx?batch={body['batch_id']}")
        self.assertEqual(r.status_code, 200)
        self.assertIn("batch-", r["Content-Disposition"])
        docs = list(load_workbook(io.BytesIO(r.content))["Documents"].iter_rows(values_only=True))
        rows = {row[0]: dict(zip(docs[0], row)) for row in docs[1:]}
        self.assertEqual(rows["good.pdf"]["Status"], "Done")
        self.assertEqual(rows["bad.pdf"]["Status"], "Failed")
        self.assertIn("could not read", rows["bad.pdf"]["Problem"].lower())
        self.assertEqual(self.client.get("/api/export.xlsx?batch=nonsense").status_code, 400)

    @override_settings(MAX_BATCH_FILES=2)
    def test_batch_size_limit(self):
        batch_id = self.upload(("a.pdf", pdf()), ("b.pdf", pdf())).json()["batch_id"]
        from django.core.files.uploadedfile import SimpleUploadedFile

        r = self.client.post("/api/documents", {"model": "prebuilt-invoice", "fields": ["VendorName"], "batch": batch_id,
                                                "files": [SimpleUploadedFile("c.pdf", pdf())]})
        self.assertEqual(r.status_code, 400)
        self.assertIn("at most 2", r.json()["error"])
        self.assertEqual(r.json()["rejected"][0]["code"], "batch")


class ImagesTests(WebTestCase):
    def setUp(self):
        super().setUp()
        self.sign_in()

    def _vendor(self, doc_id):
        return self.client.get(f"/api/documents/{doc_id}").json()["view"]["fields"][0]

    def test_text_only_by_default_and_images_on_request(self):
        logo = self.sample_bytes("inv-09-logo-brand")
        off = self.upload(("logo.pdf", logo), fields=["VendorName"]).json()["accepted"][0]["id"]
        on = self.client.post("/api/documents", {"model": "prebuilt-invoice", "fields": ["VendorName"], "include_images": "1",
                                                 "files": [SimpleUploadedFile("logo.pdf", logo)]}).json()["accepted"][0]["id"]
        self.wait()
        self.assertEqual(self._vendor(off)["value"], "Shreeji Novaprint Solutions Pvt. Ltd.")
        self.assertEqual(self._vendor(on)["value"], "Novaprint")
        page = self.client.get(f"/partials/review/{off}")
        self.assertContains(page, "Typed text only:")
        self.assertContains(page, "1 image (logos, stamps or photos) was left out")
        self.assertNotContains(page, "Pictures <span")
        page = self.client.get(f"/partials/review/{on}")
        self.assertContains(page, "Images included:")
        self.assertContains(page, f"/api/documents/{on}/figures/1.1")

    def test_picture_download_and_zip(self):
        import zipfile

        r = self.client.post("/api/samples/id-01-specimen-card", {"include_images": "1",
                                                                  "fields": ["FirstName", "LastName"]})
        self.assertEqual(r.json()["model_id"], "prebuilt-idDocument")
        doc = r.json()["accepted"][0]["id"]
        self.wait()
        png = self.client.get(f"/api/documents/{doc}/figures/1.1")
        self.assertEqual(png["Content-Type"], "image/png")
        self.assertTrue(png.content.startswith(b"\x89PNG"))
        self.assertIn("attachment", self.client.get(f"/api/documents/{doc}/figures/1.1?download=1")["Content-Disposition"])
        self.assertEqual(self.client.get(f"/api/documents/{doc}/figures/9.9").status_code, 404)
        self.assertEqual(self.client.get(f"/api/documents/{doc}/figures/..%2Fx").status_code, 404)
        z = self.client.get(f"/api/documents/{doc}/pictures.zip")
        names = zipfile.ZipFile(io.BytesIO(z.content)).namelist()
        self.assertEqual(len(names), 1)
        self.assertTrue(names[0].endswith("/picture-1.1.png"))
        batch = self.client.get(f"/api/export-pictures.zip?batch={r.json()['batch_id']}")
        self.assertEqual(batch.status_code, 200)

    def test_no_pictures_without_images(self):
        doc = self.upload(("a.pdf", self.sample_bytes("inv-02-us-logistics"))).json()["accepted"][0]["id"]
        self.wait()
        self.assertEqual(self.client.get(f"/api/documents/{doc}/pictures.zip").status_code, 404)
        self.assertEqual(self.client.get(f"/api/documents/{doc}/figures/1.1").status_code, 404)

    def test_images_count_pages_twice(self):
        from docintel.pages import budget_usage

        self.client.post("/api/documents", {"model": "prebuilt-invoice", "fields": ["VendorName"], "include_images": "1",
                                            "files": [SimpleUploadedFile("two.pdf", pdf(2))]})
        self.assertEqual(budget_usage(200)["used"], 4)
