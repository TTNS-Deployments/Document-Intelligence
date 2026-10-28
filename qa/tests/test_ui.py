"""UI tests with Playwright (headless Chromium) against the web app in local mode. IDs match qa/test_cases.md."""

import csv
import io
import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import PASSWORD, ROOT, sample_path, sign_in

AXE = (ROOT / "qa/vendor/axe-4.11.0.min.js").read_text()


def run_sample(page, sample="inv-01-gst-software", preset=None):
    if preset:
        page.select_option("#preset-select", label=preset)
    page.select_option("#sample-select", sample)
    page.click("#run-sample")
    wait_for_review(page, sample)


def wait_for_review(page, name):
    """The finished review of this file (an earlier review stays on screen until htmx swaps in the new one)."""
    page.wait_for_selector(f'[data-review]:has(.ss-name:text-matches("^{name}\\.")) [data-page-wrap][data-rendered]')


def save(download):
    """openpyxl needs the file extension, which Playwright's temporary download path lacks."""
    import tempfile

    path = Path(tempfile.mkdtemp()) / download.suggested_filename
    download.save_as(path)
    return path


def upload(page, *paths):
    page.set_input_files("#file-input", list(paths))


# ------------------------------------------------------------------------------------------- sign-in and roles

def test_U01_sign_in_as_admin_and_user(page):
    sign_in(page, "Admin")
    assert page.locator(".role-badge").inner_text() == "Admin"
    assert page.locator(".mainnav a", has_text="Settings").count() == 1
    page.click("text=Sign out")
    page.wait_for_url("**/auth/login")
    sign_in(page, "User")
    assert page.locator(".role-badge").inner_text() == "User"
    assert page.locator(".mainnav a", has_text="Settings").count() == 0
    response = page.goto("/settings")
    assert response.status == 403 and page.locator("h1", has_text="No access").count() == 1


def test_U02_wrong_password_message(page, request):
    request.node.allow_console_errors = True  # the 403 response is logged by the browser
    sign_in(page, "Admin", password="wrong")
    assert page.locator(".msg-error", has_text="Wrong password").is_visible()
    assert "/auth/login" in page.url


# ------------------------------------------------------------------------------------------- step 1

def test_U03_model_preset_and_fields(page):
    sign_in(page)
    options = page.locator("#model-select option")
    assert options.count() == 4
    card = page.locator("#model-select option[value=prebuilt-businessCard]")
    assert card.is_disabled() and "not available" in card.inner_text()
    assert page.locator("#preset-select").input_value() == "preset:ap-basics"
    assert page.locator("#field-count").inner_text() == "9 fields selected"
    page.select_option("#preset-select", label="Tax review")
    assert page.locator("#field-count").inner_text() == "9 fields selected"
    assert page.locator('#picker input[value="TaxDetails"]').is_checked()
    page.check('#picker input[value="KVKNumber"]')
    assert page.locator("#preset-select").input_value() == "custom"
    page.fill("#picker input[type=search]", "address")
    visible = page.locator("#picker .picker-field:visible")
    assert visible.count() == 12  # the address fields of every group, and the whole "Addresses" group
    page.click("#picker button[data-act=clear]")  # clears only what the search shows
    page.fill("#picker input[type=search]", "")
    page.click("#picker button[data-act=all]")
    assert page.locator("#field-count").inner_text() == "35 fields selected"
    page.select_option("#model-select", "prebuilt-receipt")
    assert page.locator("#field-count").inner_text() == "6 fields selected"


def test_U04_personal_selection_is_remembered(page):
    sign_in(page)
    page.select_option("#preset-select", label="Tax review")
    page.click("#remember-selection")
    page.reload()
    assert page.locator("#preset-select").input_value() == "mine"
    assert page.locator('#picker input[value="TaxDetails"]').is_checked()


# ------------------------------------------------------------------------------------------- upload and batch

def test_U05_upload_one_file_and_review(page):
    sign_in(page)
    upload(page, sample_path("inv-03-uk-interiors"))
    page.wait_for_selector(".batch-item .chip-done")
    page.wait_for_selector("[data-review] [data-page-wrap][data-rendered]")
    assert page.locator(".ss-name").inner_text() == "inv-03-uk-interiors.pdf"
    assert "API 2024-11-30" in page.locator(".ss-meta").inner_text()
    assert page.locator(".stepper li.is-current").inner_text().strip().endswith("Review")


def test_U06_batch_statuses_update_live_and_one_excel_at_the_end(page):
    sign_in(page)
    files = [sample_path("inv-02-us-logistics"), sample_path("inv-06-au-consulting"),
             sample_path("inv-07-sg-discount", "png"), sample_path("inv-04-eu-energy")]
    with page.expect_download(timeout=30000) as dl:
        upload(page, *files)
        page.locator(".batch-item").first.wait_for()
        assert page.locator("#auto-download").is_checked()
        page.locator(".batch-export button[disabled]", has_text="Waiting").wait_for()  # locked while documents run
    # the Excel file downloads by itself once, when every document has finished
    assert dl.value.suggested_filename.startswith("batch-") and dl.value.suggested_filename.endswith(".xlsx")
    wb = load_workbook(save(dl.value))
    rows = list(wb["Documents"].iter_rows(values_only=True))
    assert len(rows) == 5 and {r[4] for r in rows[1:]} == {"Done"}
    assert page.locator(".batch-head", has_text="4 of 4 finished").is_visible()
    assert page.locator(".batch-export.is-ready", has_text="All 4 documents are finished").is_visible()
    assert page.locator(".batch-item .chip-done").count() == 4
    with page.expect_download() as again:
        page.click("[data-batch-excel]")
    assert load_workbook(save(again.value))["Documents"].max_row == 5  # the button gives the same file


def test_U07_rejected_files_show_reasons(page, request):
    request.node.allow_console_errors = True
    sign_in(page)
    bad = Path(page.context.browser.version and "/tmp") / "notes.txt"
    bad.write_text("hello")
    fake_pdf = Path("/tmp") / "fake.pdf"
    fake_pdf.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 64)
    upload(page, str(bad), str(fake_pdf))
    page.wait_for_selector(".rejection >> text=does not match")
    texts = page.locator(".rejection").all_inner_texts()
    assert any("only PDF, JPEG and PNG" in t for t in texts)


# ------------------------------------------------------------------------------------------- review

def _sidecar(sample_id):
    return json.loads((ROOT / f"Data/sample_invoices/{sample_id}.json").read_text())


def _check_boxes(page, sample_id, variant, names):
    from docintel import samples as samples_mod

    sample = _sidecar(sample_id)
    wrap = page.locator("[data-page-wrap]").bounding_box()
    page_w = sample["pages"][0]["width"] * (sample["variants"][variant].get("scale", 1) if variant == "png" else 1)
    checked = 0
    for name in names:
        spec = sample["fields"][name]
        mapped = samples_mod.variant_polygon(sample, variant, spec["polygon"], spec["page"])
        _, poly = mapped
        xs, ys = poly[0::2], poly[1::2]
        scale = wrap["width"] / page_w
        expected = (wrap["x"] + min(xs) * scale, wrap["y"] + min(ys) * scale,
                    wrap["x"] + max(xs) * scale, wrap["y"] + max(ys) * scale)
        box = page.locator(f'polygon[data-path="{name}"]').first.bounding_box()
        got = (box["x"], box["y"], box["x"] + box["width"], box["y"] + box["height"])
        assert all(abs(a - b) <= 1.5 for a, b in zip(got, expected)), (name, got, expected)
        checked += 1
    return checked


def test_U08_overlay_boxes_sit_on_the_sidecar_polygons_pdf(page):
    sign_in(page)
    run_sample(page, "inv-01-gst-software", preset="Full invoice with line items")
    assert _check_boxes(page, "inv-01-gst-software", "pdf",
                        ["VendorName", "InvoiceId", "InvoiceDate", "SubTotal", "InvoiceTotal", "CustomerTaxId"]) == 6
    page.click("[data-zoom-in]")
    page.wait_for_function("document.querySelector('[data-zoom-level]').textContent === '125%'")
    assert _check_boxes(page, "inv-01-gst-software", "pdf", ["InvoiceTotal"]) == 1


def test_U09_overlay_boxes_on_rotated_scan_png(page):
    sign_in(page)
    run_sample(page, "inv-08-scanned-print")
    assert page.locator(".file-badge").last.inner_text() == "PNG"
    assert _check_boxes(page, "inv-08-scanned-print", "png", ["VendorName", "InvoiceId", "AmountDue"]) == 3


def test_U10_hover_and_click_link_table_and_boxes(page):
    sign_in(page)
    run_sample(page, "inv-04-eu-energy")
    page.hover('tr.field-row[data-path="InvoiceTotal"]')
    assert "is-active" in page.locator('polygon[data-path="InvoiceTotal"]').get_attribute("class")
    page.locator('polygon[data-path="VendorName"]').click()
    assert "is-active" in page.locator('tr.field-row[data-path="VendorName"]').get_attribute("class")
    page.mouse.move(5, 5)
    assert "is-active" in page.locator('tr.field-row[data-path="VendorName"]').get_attribute("class")  # stays selected


def test_U11_multi_page_navigation_follows_the_field(page):
    sign_in(page)
    run_sample(page, "inv-05-multipage-supplies", preset="Full invoice with line items")
    assert page.locator("[data-page-count]").inner_text() == "2"
    page.hover('tr.field-row[data-path="AmountDue"]')  # printed on page 2
    page.wait_for_function("document.querySelector('[data-page-num]').textContent === '2'")
    page.click("[data-page-prev]")
    page.wait_for_function("document.querySelector('[data-page-num]').textContent === '1'")
    items = page.locator("section.items", has=page.locator("h3", has_text="Line items"))
    assert items.locator("tbody tr").count() == 42
    assert items.locator("td.is-review").count() == 2


def test_U12_not_found_and_review_chips(page):
    sign_in(page)
    run_sample(page, "inv-01-gst-software", preset="Header with addresses")
    assert page.locator('tr[data-path="ShippingAddress"] .chip-missing', has_text="Not found").is_visible()
    assert page.locator('tr[data-path="PurchaseOrder"] .chip-review', has_text="Review").is_visible()
    assert page.locator("tr.field-row.is-review").count() == 1
    assert re.search(r"1 of 13\s+fields need review", page.locator(".ss-review").inner_text())


def test_U13_change_fields_without_new_analysis(page):
    sign_in(page)
    run_sample(page, "inv-02-us-logistics")
    page.click("text=Change fields")
    dialog = page.locator("dialog[open]")
    dialog.locator('input[value="ShippingAddress"]').check()
    dialog.locator('input[value="DueDate"]').uncheck()
    dialog.locator("button[type=submit]").click()
    page.wait_for_selector('tr.field-row[data-path="ShippingAddress"]')
    assert page.locator('tr.field-row[data-path="DueDate"]').count() == 0
    assert page.locator('polygon[data-path="ShippingAddress"]').count() == 1  # boxes redrawn


def test_U14_inline_edit_and_save(page):
    sign_in(page)
    run_sample(page, "inv-01-gst-software", preset="Full invoice with line items")
    page.click('tr[data-path="PurchaseOrder"] [data-edit]')
    dialog = page.locator("dialog[open]")
    assert dialog.locator("[data-edit-model]").inner_text() == "PO-MRV-7781"
    dialog.locator("[data-edit-value]").fill("PO-MRV-7782")
    dialog.locator("button[type=submit]").click()
    row = page.locator('tr[data-path="PurchaseOrder"]')
    row.locator(".chip-edited").wait_for()
    assert "PO-MRV-7782" in row.inner_text() and "Model value: PO-MRV-7781" in row.inner_text()
    assert re.search(r"1 of 19\s+fields need review", page.locator(".ss-review").inner_text())  # CustomerTaxId left
    page.dblclick('.items-table td[data-path="Items[0].Amount"]')
    page.locator("dialog[open] [data-edit-value]").fill("84500")
    page.locator("dialog[open] button[type=submit]").click()
    page.wait_for_selector('.items-table td.is-edited[data-path="Items[0].Amount"]')
    page.click("text=Mark all reviewed")
    page.wait_for_selector(".ss-review >> text=No fields need review")


def test_U15_exports_download_and_open(page):
    sign_in(page)
    run_sample(page, "inv-03-uk-interiors", preset="Full invoice with line items")
    for fmt in ("XLSX", "CSV", "JSON"):
        page.click("[data-menu-button]")
        with page.expect_download() as dl:
            page.click(f".menu-list a:text('Download {fmt}')")
        path = save(dl.value)
        assert dl.value.suggested_filename == f"inv-03-uk-interiors-fields.{fmt.lower()}"
        if fmt == "XLSX":
            wb = load_workbook(path)
            assert wb.sheetnames == ["Documents", "Line items", "Details"]
            docs = list(wb["Documents"].iter_rows(values_only=True))
            assert dict(zip(docs[0], docs[1]))["Invoice total"] == 6597.6
        elif fmt == "CSV":
            rows = list(csv.reader(io.StringIO(Path(path).read_text(encoding="utf-8-sig"))))
            assert dict(zip(rows[0], rows[1]))["Vendor name"] == "Maple & Larch Interiors Ltd"
        else:
            data = json.loads(Path(path).read_text())
            assert data["fields"]["InvoiceId"]["value"] == "ML-2026-0371" and data["api_version"] == "2024-11-30"
    page.click("[data-menu-button]")
    page.click("text=Copy JSON")
    page.wait_for_selector(".toast >> text=copied")
    assert '"model": "prebuilt-invoice"' in page.evaluate("navigator.clipboard.readText()")


def test_U16_failed_document_shows_reason_and_retry(page, tmp_path):
    from reportlab.pdfgen import canvas

    path = tmp_path / "unreadable.pdf"
    c = canvas.Canvas(str(path))
    c.drawString(72, 720, "DOCINTEL-MOCK-FAIL")
    c.save()
    sign_in(page)
    upload(page, str(path))
    page.wait_for_selector(".failure >> text=could not be analysed")
    assert page.locator(".batch-item .chip-failed").is_visible()
    page.locator(".failure [data-retry]").click()
    page.wait_for_selector(".toast >> text=Queued again")
    page.wait_for_selector(".failure >> text=could not be analysed")


# ------------------------------------------------------------------------------------------- history

def test_U17_history_filters_open_and_delete(page):
    sign_in(page)
    run_sample(page, "inv-06-au-consulting")
    page.goto("/history")
    row = page.locator("tr", has_text="inv-06-au-consulting.pdf").first
    assert row.is_visible()
    page.fill("#history-q", "no-such-document")
    page.wait_for_selector("text=No documents match")
    assert "q=no-such-document" in page.url
    page.fill("#history-q", "au-consulting")
    page.wait_for_selector("tr:has-text('inv-06-au-consulting.pdf')")
    page.check("input[name=mine]")
    page.wait_for_url("**mine=1**")
    page.wait_for_selector("tr:has-text('inv-06-au-consulting.pdf')")
    page.locator("[data-check]").first.check()
    assert page.locator("#bulk-bar").is_visible()
    with page.expect_download() as dl:
        page.click("[data-bulk-export=csv]")
    assert dl.value.suggested_filename.startswith("documents-")
    count = page.locator("tbody tr").count()
    page.locator("tbody tr [data-delete]").first.click()
    dialog = page.locator("dialog[open]")
    assert "cannot be undone" in dialog.inner_text()
    dialog.locator("button", has_text="Cancel").click()
    assert page.locator("tbody tr").count() == count
    page.locator("tbody tr [data-delete]").first.click()
    page.locator("dialog[open] button", has_text="Delete").click()
    page.wait_for_function(f"document.querySelectorAll('tbody tr').length === {count - 1}")


def test_U18_admin_manages_presets(page):
    sign_in(page, "Admin")
    page.goto("/settings")
    page.click("#preset-new")
    page.fill("#preset-name", "QA quick check")
    page.check('#preset-picker input[value="VendorName"]')
    page.check('#preset-picker input[value="AmountDue"]')
    page.click("#preset-form button[type=submit]")
    page.wait_for_selector(".preset-item >> text=QA quick check")
    page.goto("/")
    page.select_option("#preset-select", label="QA quick check")
    assert page.locator("#field-count").inner_text() == "2 fields selected"
    page.goto("/settings")
    page.locator(".preset-item", has_text="QA quick check").locator("[data-delete-preset]").click()
    page.locator("dialog[open] button", has_text="Delete").click()
    page.wait_for_selector(".preset-item >> text=QA quick check", state="detached")
    page.fill("#review_threshold", "0.9")
    page.click("text=Save settings")
    assert page.locator(".msg-success").is_visible()
    page.fill("#review_threshold", "0.8")
    page.click("text=Save settings")


# ------------------------------------------------------------------------------------------- front end quality

def test_U19_theme_toggle(page):
    sign_in(page)
    before = page.evaluate("document.documentElement.dataset.theme")
    page.click("[data-theme-toggle]")
    after = page.evaluate("document.documentElement.dataset.theme")
    assert {before, after} == {"light", "dark"}
    page.reload()
    assert page.evaluate("document.documentElement.dataset.theme") == after
    bg = page.evaluate("getComputedStyle(document.body).backgroundColor")
    assert bg == ("rgb(13, 17, 23)" if after == "dark" else "rgb(245, 246, 248)")


def test_U20_keyboard_navigation(page):
    sign_in(page)
    page.keyboard.press("Tab")
    assert page.evaluate("document.activeElement.className") == "skip-link"
    # Enter on the upload area opens the file picker. Headless Chromium reports the native dialog only some of the
    # time, so the test checks that the hidden file input was activated, then supplies the file.
    page.evaluate("document.getElementById('file-input').addEventListener('click', () => window.__picker = true)")
    page.focus("#dropzone")
    page.keyboard.press("Enter")
    assert page.evaluate("window.__picker === true")
    page.set_input_files("#file-input", sample_path("inv-02-us-logistics"))
    page.wait_for_selector("[data-review] [data-page-wrap][data-rendered]")
    first = page.locator("tr.field-row").first
    first.focus()
    page.keyboard.press("ArrowDown")
    second_path = page.locator("tr.field-row").nth(1).get_attribute("data-path")
    assert page.evaluate("document.activeElement.dataset.path") == second_path
    page.keyboard.press("Enter")
    assert "is-active" in page.locator(f'polygon[data-path="{second_path}"]').get_attribute("class")
    outline = page.evaluate("getComputedStyle(document.activeElement, ':focus-visible').outlineStyle")
    assert outline != "none"


@pytest.mark.parametrize("path", ["/auth/login", "/", "/history", "/settings", "review"])
def test_U21_no_horizontal_overflow_at_375px(page, path):
    page.set_viewport_size({"width": 375, "height": 800})
    if path != "/auth/login":
        sign_in(page, "Admin")
    if path == "review":
        run_sample(page, "inv-01-gst-software", preset="Full invoice with line items")
    else:
        page.goto(path)
    overflow = page.evaluate("document.documentElement.scrollWidth - document.documentElement.clientWidth")
    assert overflow <= 0, f"{path} overflows by {overflow}px"


@pytest.mark.parametrize("path", ["/auth/login", "/", "/history", "/settings", "review"])
def test_U22_accessibility_sanity(page, path):
    """axe-core: no serious or critical violations (WCAG 2.1 A/AA rules)."""
    if path != "/auth/login":
        sign_in(page, "Admin")
    if path == "review":
        run_sample(page, "inv-01-gst-software", preset="Full invoice with line items")
    else:
        page.goto(path)
    page.add_script_tag(content=AXE)
    result = page.evaluate("""async () => {
        const r = await axe.run(document, {runOnly: {type: 'tag', values: ['wcag2a', 'wcag2aa', 'wcag21a', 'wcag21aa']}});
        return r.violations.filter(v => ['serious', 'critical'].includes(v.impact))
                 .map(v => v.id + ': ' + v.nodes.slice(0, 3).map(n => n.target.join(' ')).join(' | '));
    }""")
    assert not result, result


@pytest.mark.parametrize("theme", ["light", "dark"])
def test_U23_contrast_in_both_themes(page, theme):
    sign_in(page)
    page.evaluate(f"localStorage.setItem('docintel.theme', '{theme}')")
    run_sample(page, "inv-01-gst-software", preset="Header with addresses")
    page.add_script_tag(content=AXE)
    result = page.evaluate("""async () => (await axe.run(document, {runOnly: ['color-contrast']}))
        .violations.map(v => v.nodes.slice(0, 3).map(n => n.target.join(' ')).join(' | '))""")
    assert not result, result


def test_U24_reduced_motion_is_respected(browser, server):
    context = browser.new_context(base_url=server, reduced_motion="reduce")
    page = context.new_page()
    sign_in(page)
    duration = page.evaluate("getComputedStyle(document.querySelector('.btn')).transitionDuration")
    context.close()
    assert duration in ("1e-06s", "0.001ms", "1e-6s") or float(duration.rstrip("ms")) < 0.01, duration


def test_U25_main_flow_with_csp_enforced(browser, server):
    """Same browser, CSP not bypassed: the app's own scripts, pdf.js and its worker must all run under the CSP."""
    context = browser.new_context(base_url=server, viewport={"width": 1280, "height": 900})
    page = context.new_page()
    errors = []
    page.on("console", lambda m: errors.append(m.text) if m.type == "error" else None)
    page.on("pageerror", lambda e: errors.append(str(e)))
    sign_in(page, "Admin")
    run_sample(page, "inv-02-us-logistics")
    assert page.locator("polygon").count() > 3
    page.goto("/history")
    page.goto("/settings")
    page.locator(".preset-item").first.wait_for()
    context.close()
    assert not [e for e in errors if "Content Security Policy" in e or "Refused" in e], errors
    assert not errors, errors


# ------------------------------------------------------------------------------------------- images

def test_U26_include_images_switch(page):
    sign_in(page)
    switch = page.locator("#include-images")
    assert not switch.is_checked()
    assert page.locator("[data-images-off]").is_visible() and not page.locator("[data-images-on]").is_visible()
    run_sample(page, "inv-09-logo-brand")
    row = page.locator('tr.field-row[data-path="VendorName"]')
    assert "Shreeji Novaprint Solutions Pvt. Ltd." in row.inner_text()  # typed name, not the logo
    assert "Typed text only" in page.locator(".images-note").inner_text()
    assert page.locator(".pictures").count() == 0

    page.locator(".switch").click()  # the visible switch toggles the hidden checkbox
    assert switch.is_checked() and page.locator("[data-images-on]").is_visible()
    page.click("#run-sample")
    page.wait_for_selector('.review .images-note:has-text("Images included")')
    page.wait_for_selector("[data-review] [data-page-wrap][data-rendered]")
    assert "Novaprint" in page.locator('tr.field-row[data-path="VendorName"] .value-text').inner_text()
    picture = page.locator(".picture").first
    picture.locator("img").wait_for()
    page.wait_for_function("[...document.querySelectorAll('.picture img')].every(i => i.complete && i.naturalWidth > 0)")
    picture.hover()
    assert "is-active" in page.locator('polygon[data-path="picture-1.1"]').get_attribute("class")
    with page.expect_download() as dl:
        picture.locator("a[download]").click()
    assert dl.value.suggested_filename.endswith("-picture-1.1.png")


def test_U27_id_card_photo_is_extracted(page):
    sign_in(page)
    page.locator(".switch").click()
    page.select_option("#sample-select", "id-01-specimen-card")
    page.click("#run-sample")
    page.wait_for_selector('[data-review]:has(.ss-name:text-matches("^id-01-specimen-card\\.")) [data-page-wrap][data-rendered]')
    assert page.locator("#model-select").input_value() == "prebuilt-idDocument"
    assert "RAO" in page.locator('tr.field-row[data-path="LastName"]').inner_text()
    page.wait_for_function("document.querySelector('.picture img') && document.querySelector('.picture img').naturalWidth > 0")
    assert page.locator(".picture").count() == 1
    assert page.locator("polygon.is-picture").count() == 1
